"""
Convert FINAL_REPORT.md to Word Document (High Detail Version)
"""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_code_block(doc, code_text):
    """Adds a code block style paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    runner = p.add_run(code_text)
    runner.font.name = 'Consolas'
    runner.font.size = Pt(9)
    runner.font.color.rgb = RGBColor(50, 50, 50)
    
    # Add a border (requires easy XML hack or just simple formatting)
    # Keeping it simple for stability: Indented Consolas text

def create_word_report():
    doc = Document()
    
    # ---------------------------------------------------------
    # Styles
    # ---------------------------------------------------------
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # ---------------------------------------------------------
    # Title Page
    # ---------------------------------------------------------
    title = doc.add_heading('IoTGuard - Final Project Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('ML-Driven IoT Intrusion Detection System')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.italic = True
    
    doc.add_paragraph()
    
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = True
    tbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell_1 = tbl.cell(0, 0)
    cell_1.text = "Prepared For:\nProject Evaluation Committee"
    cell_2 = tbl.cell(0, 1)
    cell_2.text = "Prepared By:\nIoTGuard Team"
    
    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run('Project Repository: ').bold = True
    info.add_run('https://github.com/A7medico/iotguard\n')
    info.add_run('Version: ').bold = True
    info.add_run('1.0.0\n')
    info.add_run('Date: ').bold = True
    info.add_run('December 2025')
    
    doc.add_page_break()
    
    # ---------------------------------------------------------
    # 1. Executive Summary
    # ---------------------------------------------------------
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        'IoTGuard is a lightweight, machine learning-driven Intrusion Detection System (IDS) '
        'specifically engineered for the resource-constrained and heterogeneous nature of Internet of Things (IoT) networks. '
        'Unlike traditional IT security solutions that rely on heavy signature databases, IoTGuard utilizes behavioral analysis '
        'through advanced ensemble machine learning to detect both known attacks (e.g., Mirai botnet, SYN floods) and '
        'zero-day anomalies in real time.'
    )
    doc.add_paragraph(
        'The system addresses the critical security gap in the rapidly expanding IoT ecosystem, where devices are often '
        'shipped with vulnerabilities and lack support for endpoint protection agents. IoTGuard operates at the network level, '
        'providing a universal security layer that monitors, analyzes, and actively protects all connected devices regardless '
        'of their make, model, or operating system.'
    )
    
    doc.add_heading('Key Achievements', level=2)
    achievements = [
        ('Real-time Low-Latency Detection', 'Achieved sub-second inference latency, neutralizes threats before damage occurs.'),
        ('Advanced Multi-Model Ensemble', 'Fuses LightGBM (Precision) and IsolationForest (Anomaly) for optimal coverage.'),
        ('Precision Performance', '88.86% accuracy on complex IoT attacks, 99.47% on IT traffic, ROC-AUC > 0.97.'),
        ('Automated Active Defense', 'Cross-platform blocking (Windows/Linux) translates ML decisions into firewall rules.'),
        ('Explainable AI (XAI)', 'Integrates SHAP to provide transparent, human-readable reasons for decisions.'),
        ('Enterprise-Grade Engineering', 'Docker containerization, JWT auth, and Prometheus metrics.'),
    ]
    for title, desc in achievements:
        p = doc.add_paragraph()
        p.style = 'List Bullet'
        p.add_run(title).bold = True
        p.add_run(f': {desc}')

    doc.add_page_break()

    # ---------------------------------------------------------
    # 2. Final System Architecture
    # ---------------------------------------------------------
    doc.add_heading('2. Final System Architecture and Design', level=1)
    doc.add_paragraph(
        'The architecture of IoTGuard follows a decoupled microservices-like pattern where data ingestion, '
        'processing, decision-making, and response are separated into distinct logical components.'
    )
    
    doc.add_heading('2.1 High-Level Component Diagram', level=2)
    doc.add_paragraph(
        'The following PlantUML code describes the system architecture. '
        'You can render this utilizing any PlantUML server or plugin.'
    )
    
    # PLANTUML BLOCK 1
    uml_arch = """@startuml
!theme plain
package "Network Layer" {
  [IoT Devices] as IoT
  [IT Infrastructure] as IT
}
package "IoTGuard Core System" {
  component "Suricata IDS" as Suricata
  package "Decision Engine" {
    component "Ensemble Manager" as Ensemble
    component "Decision Loop" as Decision
  }
  package "Response" {
    component "Blocker" as Blocker
    component "Alerting Service" as Alerting
  }
}
IoT --> Suricata : Traffic
Suricata --> Ensemble : Logs
Ensemble --> Decision : Scores
Decision --> Blocker : Block
Decision --> Alerting : Alert
@enduml"""
    add_code_block(doc, uml_arch)
    doc.add_paragraph('Figure 1: High-Level Component Interaction', style='Caption')

    doc.add_heading('2.2 Data Flow Pipeline', level=2)
    doc.add_paragraph('The system operates as a continuous streaming pipeline:')
    
    steps = [
        ('Traffic Ingestion Layer', 'Suricata captures traffic and outputs JSON logs (eve.json).'),
        ('Feature Engineering Layer', 'Aggregates events into 10-second windows and computes 13 statistical features.'),
        ('Context & Fingerprinting', 'Dynamically classifies devices (IoT vs IT) to route to the correct ML model.'),
        ('Ensemble Inference Layer', 'Combines Supervised (LightGBM) and Unsupervised (IsolationForest) scores.'),
        ('Decision & Response Layer', 'Applies adaptive thresholds, triggers blocking, and dispatches alerts.'),
    ]
    for i, (title, text) in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.add_run(f'{i}. {title}: ').bold = True
        p.add_run(text)

    doc.add_heading('2.3 Feature Engineering Matrix', level=2)
    doc.add_paragraph('The system relies on 13 robust flow-based features agnostic to payload encryption.')
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Feature'
    hdr_cells[1].text = 'Type'
    hdr_cells[2].text = 'Rationale'
    
    features = [
        ('flows', 'Count', 'High variance indicates scanning/DDoS'),
        ('bytes_total', 'Volume', 'Spikes signal exfiltration or floods'),
        ('syn_ratio', 'Flag', 'High values (~1.0) confirm SYN Floods'),
        ('mean_bytes_flow', 'Stat', 'Small values typical of C2 heartbeats'),
        ('protocol_diversity', 'Meta', 'IoT devices usually have low diversity'),
        ('iat_mean', 'Time', 'Machine traffic has distinct timing vs human'),
    ]
    for f, t, r in features:
        row = table.add_row().cells
        row[0].text = f
        row[1].text = t
        row[2].text = r
    doc.add_paragraph('Table 1: Key Network Features (Subset)', style='Caption')

    doc.add_heading('2.4 Ensemble Logic', level=2)
    doc.add_paragraph('Below is the PlantUML Source for the Ensemble internal logic:')
    
    # PLANTUML BLOCK 2
    uml_logic = """@startuml
state "Ensemble Engine" {
  state ParallelProcs {
    state "Supervised (LightGBM)" as M1
    state "Unsupervised (IsoForest)" as M2
  }
  ParallelProcs --> Combine
  state Combine {
    state "Weighted Average" as Avg
    Avg : Score = (0.7 * P_attack) + (0.3 * Anomaly)
  }
}
Combine --> DecisionCheck
@enduml"""
    add_code_block(doc, uml_logic)
    doc.add_paragraph('Figure 2: Ensemble Logic Flow', style='Caption')

    doc.add_page_break()

    # ---------------------------------------------------------
    # 3. Impact
    # ---------------------------------------------------------
    doc.add_heading('3. Impact of Engineering Solutions', level=1)
    
    doc.add_heading('3.1 Global and Economic Context', level=2)
    doc.add_paragraph('The IoT market is projected to reach $3.3 trillion by 2030.')
    doc.add_paragraph('Risk Mitigation: ', style='List Bullet').add_run('IoTGuard serves as an "insurance policy" for digital transformation, lowering security barriers for developing economies.').bold = False
    doc.add_paragraph('Cost Efficiency: ', style='List Bullet').add_run('Replaces expensive appliances with software that runs on commodity hardware (e.g., Raspberry Pi).').bold = False
    
    doc.add_heading('3.2 Environmental Sustainability', level=2)
    doc.add_paragraph('Energy-Aware Security: ', style='List Bullet').add_run('Metadata analysis requires <10% CPU of deep packet inspection, saving energy in data centers.').bold = False
    doc.add_paragraph('Extending Asset Life: ', style='List Bullet').add_run('Protecting sensors from DDoS stress extends battery life and reduces e-waste.').bold = False
    
    doc.add_heading('3.3 Societal Implications', level=2)
    doc.add_paragraph('Privacy Preservation: ', style='List Bullet').add_run('Analyzes behavior (metadata), not content (payload). Respects user privacy in smart homes.').bold = False
    doc.add_paragraph('Trust in Automation: ', style='List Bullet').add_run('Explainable AI (SHAP) demystifies decisions, fostering trust in AI security systems.').bold = False

    doc.add_page_break()

    # ---------------------------------------------------------
    # 4. Contemporary Issues
    # ---------------------------------------------------------
    doc.add_heading('4. Contemporary Issues', level=1)
    
    doc.add_heading('4.1 The Asymmetric Information War', level=2)
    doc.add_paragraph(
        'Defenders must secure the entire surface; attackers need only one vulnerability. '
        'IoTGuard acts as a network-level palliative for unpatchable, insecure-by-design IoT devices.'
    )
    
    doc.add_heading('4.2 "Living off the Land" in IoT', level=2)
    doc.add_paragraph(
        'Attackers now use native tools (wget, netcat) instead of malware. '
        'Traditional AV fails here. IoTGuard\'s behavioral anomaly detection catches the *pattern* of tool usage, not just the file hash.'
    )

    doc.add_heading('4.3 Regulatory Compliance (GDPR/NIS2)', level=2)
    doc.add_paragraph(
        'IoTGuard assists in compliance by providing mandated Incident Detection '
        'capabilities and immutable audit trails (alerts.jsonl).'
    )

    doc.add_page_break()

    # ---------------------------------------------------------
    # 5. Tools
    # ---------------------------------------------------------
    doc.add_heading('5. Tools and Technologies', level=1)
    techs = [
        ('Python 3.13+', 'Unrivaled Data Science ecosystem'),
        ('Suricata', 'Superior multi-threading and JSON logs vs Snort'),
        ('LightGBM', 'Faster training/inference than XGBoost for tabular data'),
        ('IsolationForest', 'Scales linearly for high-dimensional anomaly detection'),
        ('Docker', 'Ensures reproducible deployments'),
    ]
    for t, r in techs:
        p = doc.add_paragraph()
        p.style = 'List Bullet'
        p.add_run(t).bold = True
        p.add_run(f': {r}')
        
    doc.add_page_break()

    # ---------------------------------------------------------
    # 6. Resources
    # ---------------------------------------------------------
    doc.add_heading('6. Library and Internet Resources', level=1)
    doc.add_paragraph('Datasets:', style='List Bullet').add_run(' IoT-23 (Stratosphere Labs), CIC-IoT-2023 (DDoS testing).').bold = False
    doc.add_paragraph('Key Libraries:', style='List Bullet').add_run(' scikit-learn, shap, pandas, joblib.').bold = False

    doc.add_page_break()

    # ---------------------------------------------------------
    # 7. Test Results
    # ---------------------------------------------------------
    doc.add_heading('7. Test Results', level=1)
    
    doc.add_heading('7.1 Verification Matrix', level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text = 'Module'
    table.rows[0].cells[1].text = 'Scope'
    table.rows[0].cells[2].text = 'Result'
    
    tests = [
        ('IP Blocker', 'Command Injection Prevention', 'PASS'),
        ('Ensemble', 'Score Fusion Logic', 'PASS'),
        ('Alerting', 'Rate Limiting & Dispatch', 'PASS'),
        ('Fingerprint', 'Device Classification', 'PASS'),
        ('Integration', 'End-to-End Smoke Test', 'PASS'),
    ]
    for m, s, r in tests:
        row = table.add_row().cells
        row[0].text = m
        row[1].text = s
        row[2].text = r

    doc.add_heading('7.2 Model Performance', level=2)
    doc.add_paragraph('IoT Environment Model:', style='List Bullet')
    doc.add_paragraph('    Accuracy: 88.86%')
    doc.add_paragraph('    ROC-AUC: 0.9765')
    doc.add_paragraph('IT Environment Model:', style='List Bullet')
    doc.add_paragraph('    Accuracy: 99.47%')
    
    doc.add_page_break()

    # ---------------------------------------------------------
    # 8. Conclusion
    # ---------------------------------------------------------
    doc.add_heading('8. Conclusion and Future Work', level=1)
    doc.add_paragraph(
        'IoTGuard demonstrates that enterprise-grade security does not require enterprise-grade hardware. '
        'By utilizing efficient feature engineering and ensemble ML, we successfully defend the IoT edge.'
    )
    
    doc.add_heading('Future Roadmap', level=2)
    doc.add_paragraph('Federated Learning (Privacy-preserving collaboration)', style='List Bullet')
    doc.add_paragraph('Deep Learning (LSTM/GRU) for sequence analysis', style='List Bullet')
    doc.add_paragraph('eBPF Integration for kernel-level blocking', style='List Bullet')
    
    doc.add_paragraph()
    doc.add_paragraph('GitHub: https://github.com/A7medico/iotguard')

    # Save
    output_path = 'project_documentation/FINAL_REPORT_detailed.docx'
    doc.save(output_path)
    print(f'Word document created: {output_path}')

if __name__ == '__main__':
    create_word_report()
